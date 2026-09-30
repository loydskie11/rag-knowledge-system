import re

with open(r"c:\Projects\rag-governance\backend\main.py", "r", encoding="utf-8") as f:
    content = f.read()

# Fix the specific broken block
broken = """        elif filename_lower.endswith(".docx"):
            import docx as _docx
                        doc = docx.Document(io.BytesIO(contents))
            parts = []"""
fixed = """        elif filename_lower.endswith(".docx"):
            import docx
            doc = docx.Document(io.BytesIO(contents))
            parts = []"""
content = content.replace(broken, fixed)

# Check for other broken blocks
broken2 = """    elif filename_lower.endswith(".docx"):
        import docx
                    doc = docx.Document(io.BytesIO(contents))
            parts = []"""
fixed2 = """    elif filename_lower.endswith(".docx"):
        import docx
        doc = docx.Document(io.BytesIO(contents))
        parts = []
        for block in _iter_block_items(doc):
            if hasattr(block, "text"):
                parts.append(block.text)
            elif hasattr(block, "rows"):
                for row in block.rows:
                    for cell in row.cells:
                        parts.append(cell.text)
        extracted_text = "\\n".join(parts)"""
# We must apply dynamic fixing
def fix_indent(match):
    prefix = match.group(1)
    body = match.group(2)
    # the body should be prefixed by `prefix + "    "`
    fixed_body = ""
    for line in body.split("\n"):
        fixed_body += prefix + "    " + line.strip() + "\n"
    return prefix + "elif filename_lower.endswith('.docx'):\n" + prefix + "    import docx\n" + fixed_body

# Wait, let's just do a blanket find and replace for the exact blocks instead of regex.
